"""
validate_files.py — post-generation self-check for the NDF 3-test-files skill.

Re-implements the checkable rules from references/validation_rules.md (and the
cross-file consistency checklist) as code, so step 7 of SKILL.md ("self-check
against the checklist") is a script run instead of manual re-reading. This is
NOT a substitute for reading validation_rules.md — some rules (e.g. "does the
data actually exercise the scenario the user described") need judgement and
can't be checked mechanically. But every rule with a concrete, testable
condition is checked here.

Usage:
    python validate_files.py <input_csv_path> <spec_grouping_xlsx_path> <type_grouping_xlsx_path> [<discontinued_csv_path>]

The optional 4th argument (DDB mode only) also validates the Discontinued CSV
against the Input CSV (validation_rules.md §4).

Exits non-zero and prints ERROR/WARN lines if anything looks off. Claude should
run this after build_files.build(...) and fix any ERRORs before calling
present_files — and should surface unresolved WARNs to the user rather than
silently deciding they're fine.
"""
import csv
import re
import sys
from collections import Counter

ICONS = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
ICON_SET = set(ICONS)
VALID_SUBSIDIARY = {"KOR", "MJP", "USA"}
VALID_SPEC_TYPE = {"fixed", "irregular", "range", "formula"}

# --- calculationFormula syntax rules (validation_rules.md §2) ---------------

_CF_ALLOWED_OP_CHARS = set("<>=!*+-%/")
_CF_KEYWORDS = ("RULE", "AND", "OR")
_CF_FULLWIDTH_SPACE = "\u3000"

# Each pattern matches one full statement (including its trailing ';').
_CF_PATTERNS = [
    ("CEILING", re.compile(r"^IF \[.+\] RULE CEILING\(spec_name_common, ?\d+\);$")),
    ("ROUND", re.compile(r"^IF \[.+\] RULE ROUND\(spec_name_common,\d+\);$")),
    ("tolerance_range", re.compile(r"^spec_name_common[+\-±]10%;$")),
    ("add_subtract", re.compile(r"^spec_name_common[+\-]\d+;$")),
    ("notation_equality", re.compile(r"^IF \[spec_name_common=[^\]]+\] RULE spec_name_common=[^;]+;$")),
]

# Token scanner: valid characters are word chars, whitespace, digits, brackets,
# parentheses, comma, period, underscore, ± (only meaningful in tolerance_range,
# checked separately), and the allowed single operator chars above.
_CF_INVALID_TOKEN_RE = re.compile(
    r"(==|<>|&&|\|\||[^\w\s\[\]\(\),.\-+*/%<>=!±;])"
)


def validate_calculation_formula(value):
    """Return a list of error strings for a populated calculationFormula cell.
    Call only when value is non-empty/truthy."""
    errors = []
    text = str(value)

    if not text.rstrip().endswith(";"):
        errors.append(f"must end with a semicolon (;): {text!r}")

    if _CF_FULLWIDTH_SPACE in text:
        errors.append(f"contains a full-width space; only half-width (ASCII) spaces are allowed: {text!r}")

    # Each non-blank line must itself be semicolon-terminated (one statement per line).
    lines = [ln for ln in text.split("\n")]
    for i, ln in enumerate(lines):
        if ln.strip() == "":
            continue
        if not ln.rstrip().endswith(";"):
            errors.append(f"line {i + 1} must end with ';' before the line break: {ln!r}")

    # RULE/AND/OR must have a half-width space on both sides.
    for kw in _CF_KEYWORDS:
        for m in re.finditer(re.escape(kw), text):
            start, end = m.span()
            before = text[start - 1] if start > 0 else ""
            after = text[end] if end < len(text) else ""
            if before != " " or after != " ":
                errors.append(f"'{kw}' must have a half-width space before and after it: found in {text!r}")

    # Disallowed operator tokens anywhere in the cell (e.g. ==, <>, &&, ||, or stray symbols).
    bad_tokens = _CF_INVALID_TOKEN_RE.findall(text)
    if bad_tokens:
        errors.append(f"disallowed operator/character token(s) {bad_tokens}: {text!r}")

    # '±' is only valid inside the tolerance_range pattern.
    if "±" in text:
        for stmt in [s.strip() for s in text.split("\n") if s.strip()]:
            if "±" in stmt and not _CF_PATTERNS[2][1].match(stmt):
                errors.append(f"'±' is only permitted in the tolerance-range pattern (spec_name_common±10%;): {stmt!r}")

    # Every statement (split on ';' + optional newline) must match one of the 5 permitted patterns.
    statements = [s.strip() for s in re.split(r";\s*\n?", text) if s.strip()]
    for stmt in statements:
        stmt_full = stmt if stmt.endswith(";") else stmt + ";"
        if not any(pat.match(stmt_full) for _, pat in _CF_PATTERNS):
            errors.append(f"does not match any of the 5 permitted calculationFormula patterns: {stmt_full!r}")

    return errors


class Report:
    def __init__(self):
        self.errors = []
        self.warnings = []

    def err(self, msg):
        self.errors.append(msg)

    def warn(self, msg):
        self.warnings.append(msg)

    def ok(self):
        return not self.errors

    def print_all(self):
        for e in self.errors:
            print(f"ERROR: {e}")
        for w in self.warnings:
            print(f"WARN: {w}")
        if not self.errors and not self.warnings:
            print("All checks passed, no warnings.")


def _no_bad_spacing(value):
    if value is None:
        return True
    s = str(value)
    return s == s.strip() and "  " not in s


def _icons_in(text):
    return [c for c in str(text) if c in ICON_SET]


def _params_icon_map(params_str):
    """Parse '①[a=a,b=b], ②[1-10/1]' into {icon: raw_block}, preserving the
    left-to-right order the icons were defined in (dict insertion order)."""
    out = {}
    if not params_str:
        return out
    for m in re.finditer(r"([①-⑳])\[([^\]]*)\]", params_str):
        out[m.group(1)] = m.group(2)
    return out


_RANGE_BLOCK_RE = re.compile(r"\s*(-?\d+(?:\.\d+)?)-(-?\d+(?:\.\d+)?)/(\d+(?:\.\d+)?)\s*")


def _block_tokens_both_sides(block):
    """Value tokens to check for '+'/'-' restriction: both sides of every
    key=value pair for a value-list/irregular block, or just the numeric
    min for a range block (max can't be negative if min>=0 and min<max)."""
    pairs = re.findall(r"([^,\[\]=]+)=([^,\[\]=]+)", block)
    if pairs:
        toks = []
        for k, v in pairs:
            toks.append(k.strip())
            toks.append(v.strip())
        return toks
    m = _RANGE_BLOCK_RE.fullmatch(block)
    if m:
        return [m.group(1)]  # min bound only
    return []


def _block_values_for_reachability(block):
    """Right-hand-side value tokens (the ones referenced in part_number_rules
    comparisons/assignments) for a value-list/irregular block. Returns None
    for a range block (continuous values, not enumerable) so callers can skip
    the reachability check for that icon."""
    pairs = re.findall(r"([^,\[\]=]+)=([^,\[\]=]+)", block)
    if pairs:
        return {v.strip() for k, v in pairs}
    if _RANGE_BLOCK_RE.fullmatch(block):
        return None
    return set()


def _derived_icon_assignments(rules_text):
    """Best-effort parse of 'Rule <icon>=<value>' assignments (optionally
    chained with AND) out of part_number_rules text. Returns {icon: {values}}."""
    assigned = {}
    if not rules_text:
        return assigned
    for line in rules_text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.search(r"\bRule\b(.*);\s*$", line)
        consequent = m.group(1) if m else line
        for im in re.finditer(
            r"([①-⑳])\s*=\s*(.+?)(?=\s+AND\s|\s+OR\s|;|$)", consequent
        ):
            icon, val = im.group(1), im.group(2).strip()
            assigned.setdefault(icon, set()).add(val)
    return assigned


def validate_input_csv(path, report: Report):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        raw_head = f.read(3)
        f.seek(0)
        reader = csv.DictReader(f)
        rows = list(reader)
        header = reader.fieldnames or []

    if raw_head != "\ufeff"[:0]:  # placeholder, real BOM check done on bytes below
        pass
    with open(path, "rb") as f:
        if f.read(3) != b"\xef\xbb\xbf":
            report.err(f"{path}: missing UTF-8 BOM (must be utf-8-sig)")

    required_cols = [
        "part_number", "part_number_type", "category_name_original", "category_name_common",
        "brand_name_original", "brand_name_common", "subsidary_cd", "part_number_rules",
        "params", "series_name", "cad_url", "image_alt", "image_pass", "economy_flag",
        "spec_type_1", "spec_name_original_1", "spec_name_common_1", "spec_value_1", "spec_unit_1",
    ]
    for col in required_cols:
        if col not in header:
            report.err(f"Input CSV missing required column: {col}")

    non_null_cols = [
        "part_number", "part_number_type", "category_name_original", "category_name_common",
        "brand_name_original", "brand_name_common", "subsidary_cd", "economy_flag",
        "spec_type_1", "spec_name_original_1", "spec_name_common_1", "spec_value_1", "spec_unit_1",
    ]

    seen_rows = Counter()
    categories = set()
    subsidiaries = set()
    spec_names_all = set()
    part_number_types = set()

    for i, row in enumerate(rows, start=2):  # +2: header is row1
        row_key = tuple(row.get(c, "") for c in header)
        seen_rows[row_key] += 1

        for col in non_null_cols:
            if not row.get(col):
                report.err(f"Row {i}: required column '{col}' is empty")

        for col, val in row.items():
            if not _no_bad_spacing(val):
                report.err(f"Row {i}: column '{col}' has leading/trailing/double spaces: {val!r}")

        if row.get("subsidary_cd") and row["subsidary_cd"] not in VALID_SUBSIDIARY:
            report.err(f"Row {i}: subsidary_cd {row['subsidary_cd']!r} not in {VALID_SUBSIDIARY}")
        subsidiaries.add(row.get("subsidary_cd"))
        categories.add(row.get("category_name_common"))

        for col in ("cad_url", "image_alt", "image_pass"):
            if row.get(col):
                report.warn(f"Row {i}: {col} should stay empty per §14 convention, found {row[col]!r}")

        economy_flag = row.get("economy_flag")
        if economy_flag not in ("0", "1"):
            report.err(f"Row {i}: economy_flag must be '0' or '1', got {economy_flag!r}")

        # spec_type_N / spec_value_N / params cross-checks
        params_map = _params_icon_map(row.get("params", ""))
        n = 1
        while f"spec_type_{n}" in row and row.get(f"spec_type_{n}"):
            spec_type = row[f"spec_type_{n}"]
            spec_value = row.get(f"spec_value_{n}", "")
            spec_name_common = row.get(f"spec_name_common_{n}", "")
            if spec_name_common:
                spec_names_all.add(spec_name_common)
            if spec_type not in VALID_SPEC_TYPE:
                report.err(f"Row {i}: spec_type_{n} {spec_type!r} not in {VALID_SPEC_TYPE} (note: 'choice' is never valid)")
            if spec_name_common and "/" in spec_name_common:
                report.err(f"Row {i}: spec_name_common_{n} must not contain '/': {spec_name_common!r}")
            if spec_value and ("," in spec_value or "φ" in spec_value):
                report.err(f"Row {i}: spec_value_{n} must not contain ',' or lowercase 'φ': {spec_value!r}")

            value_icons = _icons_in(spec_value)
            if spec_type == "fixed" and value_icons:
                report.err(f"Row {i}: spec_type_{n}=fixed but spec_value_{n} contains an icon: {spec_value!r}")
            if spec_type == "formula":
                if not value_icons:
                    report.err(f"Row {i}: spec_type_{n}=formula requires >=1 icon in spec_value_{n}")
                for icon in value_icons:
                    if icon not in params_map:
                        report.err(f"Row {i}: formula icon {icon} in spec_value_{n} not defined in params")
            if spec_type == "range":
                if len(value_icons) != 1:
                    report.err(f"Row {i}: spec_type_{n}=range requires exactly one icon in spec_value_{n}")
                else:
                    icon = value_icons[0]
                    block = params_map.get(icon)
                    if block is None:
                        report.err(f"Row {i}: range icon {icon} not defined in params")
                    else:
                        m = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)-(-?\d+(?:\.\d+)?)/(\d+(?:\.\d+)?)\s*", block)
                        if not m:
                            report.err(f"Row {i}: range params block for {icon} doesn't match min-max/step: {block!r}")
                        else:
                            lo, hi, step = float(m.group(1)), float(m.group(2)), float(m.group(3))
                            if not (lo < hi):
                                report.err(f"Row {i}: range for {icon}: min ({lo}) must be < max ({hi})")
                            if not (step > 0):
                                report.err(f"Row {i}: range for {icon}: step must be > 0, got {step}")
            if spec_type == "irregular":
                if len(value_icons) != 1:
                    report.err(f"Row {i}: spec_type_{n}=irregular requires exactly one icon in spec_value_{n}")
                else:
                    icon = value_icons[0]
                    block = params_map.get(icon)
                    if block is None:
                        report.err(f"Row {i}: irregular icon {icon} not defined in params")
                    elif "XYZ=XYZ" not in block:
                        report.warn(f"Row {i}: irregular params for {icon} has no XYZ=XYZ catch-all (§10 convention)")
            n += 1

        # icon consistency: part_number icons ("selectable") must be consecutive from ①, no gaps
        pn_icons = _icons_in(row.get("part_number", ""))
        if not pn_icons and row.get("params"):
            report.err(f"Row {i}: part_number has no icon but params is non-empty: {row.get('params')!r}")
        if pn_icons:
            expected = list(ICONS[: len(pn_icons)])
            if pn_icons != expected:
                report.err(f"Row {i}: part_number icons {pn_icons} are not consecutive starting at ①")
            for icon in pn_icons:
                if icon not in params_map:
                    report.err(f"Row {i}: part_number icon {icon} not defined in params")

        # selectable vs. derived icon split (test_run_conventions.md §3a)
        all_params_icons_ordered = list(params_map.keys())
        derived_icons = [ic for ic in all_params_icons_ordered if ic not in pn_icons]

        if all_params_icons_ordered:
            expected_full = list(ICONS[: len(all_params_icons_ordered)])
            if sorted(all_params_icons_ordered, key=lambda c: ICONS.index(c)) != expected_full:
                report.err(
                    f"Row {i}: params icons {all_params_icons_ordered} are not one consecutive run "
                    f"starting at ① when combining selectable + derived (§3a)"
                )
            if pn_icons:
                leading = all_params_icons_ordered[: len(pn_icons)]
                if leading != pn_icons:
                    report.err(
                        f"Row {i}: selectable icons {pn_icons} must be defined first and in order "
                        f"within params; found {leading} leading instead (§3a)"
                    )

        # +/- value restriction: selectable icon values must not contain +/-; derived icons may.
        for icon, block in params_map.items():
            if icon not in pn_icons:
                continue  # derived icon: +/- allowed freely, no check
            for tok in _block_tokens_both_sides(block):
                if "+" in tok or "-" in tok:
                    report.err(
                        f"Row {i}: selectable icon {icon} has a value containing '+'/'-' "
                        f"({tok!r}) — not allowed for selectable icons (§3a)"
                    )

        for icon in params_map:
            if icon in pn_icons:
                continue  # selectable icons are always "used" (they're the part_number suffix itself)
            if icon not in row.get("part_number_rules", ""):
                report.err(
                    f"Row {i}: derived icon {icon} is defined in params but never assigned "
                    f"a value in part_number_rules (§3a/§8)"
                )

        # part_number_rules line format + first-icon check (IF-condition vs. bare Rule assignment)
        rules_text = row.get("part_number_rules", "")
        if rules_text:
            for line in [l for l in rules_text.splitlines() if l.strip()]:
                line = line.strip()
                if not line.endswith(";"):
                    report.err(f"Row {i}: part_number_rules line doesn't end with ';': {line!r}")
                if not re.match(r"^(IF\s*\[.*\]\s*)?Rule\b", line):
                    report.err(f"Row {i}: part_number_rules line must start with optional 'IF [...]' then mandatory 'Rule': {line!r}")
                line_icons = _icons_in(line)
                has_if = bool(re.match(r"^IF\s*\[", line))
                if line_icons:
                    first_icon = line_icons[0]
                    if has_if:
                        if first_icon not in pn_icons:
                            report.err(
                                f"Row {i}: first icon in part_number_rules line's IF condition "
                                f"({first_icon}) must be a selectable icon in part_number"
                            )
                    else:
                        if first_icon not in pn_icons and first_icon not in derived_icons:
                            report.err(
                                f"Row {i}: icon {first_icon} in part_number_rules line "
                                f"not defined in params at all"
                            )

            # derived-icon reachability: every value defined for a derived icon must be
            # produced by at least one Rule branch.
            assigned = _derived_icon_assignments(rules_text)
            for icon in derived_icons:
                icon_assigned = assigned.get(icon)
                if not icon_assigned:
                    continue  # already reported as "never assigned" above; skip redundant message
                defined_values = _block_values_for_reachability(params_map.get(icon, ""))
                if defined_values is not None:
                    missing = defined_values - icon_assigned
                    if missing:
                        report.err(
                            f"Row {i}: derived icon {icon} has value(s) {sorted(missing)} defined "
                            f"in params but not reachable via any part_number_rules branch (§3a/§8)"
                        )

        part_number_types.add(row.get("part_number_type"))

    dupes = [k for k, c in seen_rows.items() if c > 1]
    if dupes:
        report.err(f"{len(dupes)} duplicate row(s) found in Input CSV")

    if len(categories - {None}) > 1:
        report.err(f"category_name_common not consistent across rows: {categories}")
    if len(subsidiaries - {None}) > 1:
        report.err(f"subsidary_cd not consistent across rows: {subsidiaries}")

    return {
        "category_name_common": next(iter(categories), None),
        "subsidary_cd": next(iter(subsidiaries), None),
        "spec_names": spec_names_all,
        "part_number_types": part_number_types,
        "brands": {row.get("brand_name_common") for row in rows},
        "rows": rows,
    }


def validate_xlsx_pair(spec_path, type_path, input_ctx, report: Report):
    from openpyxl import load_workbook

    spec_wb = load_workbook(spec_path)
    type_wb = load_workbook(type_path)

    for name in ("specNameCommon", "exception"):
        if name not in spec_wb.sheetnames:
            report.err(f"Spec Grouping file missing sheet: {name}")
    for name in ("TypeMatch", "BasicSpecDefenition"):
        if name not in type_wb.sheetnames:
            report.err(f"Type Grouping file missing sheet: {name}")

    if "exception" in spec_wb.sheetnames:
        ex_ws = spec_wb["exception"]
        if ex_ws.max_row > 1:
            report.err("exception sheet must have zero data rows (§14) but has data")

    if "specNameCommon" in spec_wb.sheetnames:
        ws = spec_wb["specNameCommon"]
        header = [c.value for c in ws[1]]
        rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2) if any(c.value for c in r)]
        priorities = Counter(r.get("priority_order") for r in rows)
        for p, c in priorities.items():
            if c > 1:
                report.err(f"specNameCommon: priority_order {p!r} not unique")
        spec_names = Counter(r.get("spec_name_common") for r in rows)
        for s, c in spec_names.items():
            if c > 1:
                report.err(f"specNameCommon: spec_name_common {s!r} not unique")
        for r in rows:
            if r.get("category_name_common") != input_ctx["category_name_common"]:
                report.err(f"specNameCommon row category_name_common {r.get('category_name_common')!r} != CSV's {input_ctx['category_name_common']!r}")
            if r.get("subsidary_cd") != input_ctx["subsidary_cd"]:
                report.err(f"specNameCommon row subsidary_cd {r.get('subsidary_cd')!r} != CSV's {input_ctx['subsidary_cd']!r}")
            if r.get("spec_name_common") not in input_ctx["spec_names"]:
                report.err(f"specNameCommon spec_name_common {r.get('spec_name_common')!r} not found in Input CSV spec_name_common_* values")
            cf = r.get("calculationFormula")
            if cf:
                for msg in validate_calculation_formula(cf):
                    report.err(f"specNameCommon calculationFormula error: {msg}")

    if "TypeMatch" in type_wb.sheetnames:
        ws = type_wb["TypeMatch"]
        header = [c.value for c in ws[1]]
        rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2) if any(c.value for c in r)]
        pairs = Counter((r.get("part_number_type"), r.get("brand_name_common")) for r in rows)
        for pair, c in pairs.items():
            if c > 1:
                report.err(f"TypeMatch: (part_number_type, brand_name_common) {pair} not unique")
        brands_used = {r.get("brand_name_common") for r in rows}
        if len(brands_used) > 1:
            report.err(f"TypeMatch should use one single fixed primary brand per §12, found: {brands_used}")
        csv_types = input_ctx["part_number_types"]
        type_types = {r.get("part_number_type") for r in rows}
        if type_types != csv_types:
            missing = csv_types - type_types
            extra = type_types - csv_types
            if missing:
                report.err(f"TypeMatch missing part_number_type(s) present in CSV: {missing}")
            if extra:
                report.err(f"TypeMatch has part_number_type(s) not present in CSV: {extra}")
        for r in rows:
            if r.get("category_name_common") != input_ctx["category_name_common"]:
                report.err(f"TypeMatch row category_name_common mismatch: {r.get('category_name_common')!r}")
            if r.get("subsidary_cd") != input_ctx["subsidary_cd"]:
                report.err(f"TypeMatch row subsidary_cd mismatch: {r.get('subsidary_cd')!r}")

    if "BasicSpecDefenition" in type_wb.sheetnames:
        ws = type_wb["BasicSpecDefenition"]
        header = [c.value for c in ws[1]]
        rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2) if any(c.value for c in r)]
        if not (1 <= len(rows) <= 2):
            report.err(f"BasicSpecDefenition should have 1-2 rows per §13, has {len(rows)}")
        names = Counter(r.get("spec_name_common") for r in rows)
        for n, c in names.items():
            if c > 1:
                report.err(f"BasicSpecDefenition: spec_name_common {n!r} not unique")
        for r in rows:
            if r.get("spec_name_common") not in input_ctx["spec_names"]:
                report.err(f"BasicSpecDefenition spec_name_common {r.get('spec_name_common')!r} not found in Input CSV")


DDB_PREFIX = "JIRA-DDB-"
DISCONTINUED_FIXED_HEADER = [
    "discontinued_part_number", "discontinued_part_number_type",
    "category_name_original_en", "category_name_original", "category_name_common",
    "brand_name_original_en", "brand_name_original", "brand_name_common",
    "subsidary_cd", "discontinued_params",
]
# discontinued column -> GDB column it is copied from (validation_rules.md §4)
_DISCONTINUED_SOURCE = {
    "discontinued_part_number_type": "part_number_type",
    "category_name_original_en": "category_name_original",
    "category_name_original": "category_name_original",
    "category_name_common": "category_name_common",
    "brand_name_original_en": "brand_name_original",
    "brand_name_original": "brand_name_original",
    "brand_name_common": "brand_name_common",
    "subsidary_cd": "subsidary_cd",
    "discontinued_params": "params",
}


def validate_discontinued(path, gdb_rows, report: Report):
    """Discontinued DB file checks (validation_rules.md §4 / test_run_conventions.md §16)."""
    with open(path, "rb") as f:
        if f.read(3) != b"\xef\xbb\xbf":
            report.err(f"{path}: missing UTF-8 BOM (must be utf-8-sig)")
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, [])
        raw_rows = list(reader)

    # header: 10 fixed columns + complete, contiguous triples for n = 1..N (N >= 1)
    fixed = len(DISCONTINUED_FIXED_HEADER)
    if header[:fixed] != DISCONTINUED_FIXED_HEADER:
        report.err(f"Discontinued header must start with the 10 fixed columns {DISCONTINUED_FIXED_HEADER}, got {header[:fixed]}")
        return
    tail = header[fixed:]
    n_slots = len(tail) // 3
    expected_tail = []
    for n in range(1, n_slots + 1):
        expected_tail += [f"recommend_part_number_{n}", f"recommend_part_number_type_{n}", f"recommend_params_{n}"]
    if n_slots < 1 or len(tail) % 3 != 0 or tail != expected_tail:
        report.err(
            f"Discontinued header after the fixed columns must be complete contiguous "
            f"recommend_* triples for n=1..N (N>=1), got {tail}"
        )
        return

    by_pn = {r.get("part_number"): r for r in gdb_rows}
    seen = Counter()
    for i, cells in enumerate(raw_rows, start=2):
        if len(cells) != len(header):
            report.err(f"Discontinued row {i}: expected {len(header)} cells, got {len(cells)}")
            continue
        row = dict(zip(header, cells))
        pn = row["discontinued_part_number"]
        seen[pn] += 1
        src = by_pn.get(pn)
        if src is None:
            report.err(f"Discontinued row {i}: discontinued_part_number {pn!r} not found in the Input CSV")
        elif not pn.startswith(DDB_PREFIX):
            report.err(f"Discontinued row {i}: {pn!r} is a normal (non-DDB) row and must not be listed as discontinued")
        if src is not None:
            for col, gdb_col in _DISCONTINUED_SOURCE.items():
                if row[col] != src.get(gdb_col, ""):
                    report.err(f"Discontinued row {i}: {col} {row[col]!r} != GDB {gdb_col} {src.get(gdb_col, '')!r} for {pn!r}")

        recs = []
        empty_seen_at = None
        for n in range(1, n_slots + 1):
            rpn = row[f"recommend_part_number_{n}"]
            rtype = row[f"recommend_part_number_type_{n}"]
            rparams = row[f"recommend_params_{n}"]
            filled = bool(rpn or rtype)
            if not filled:
                if rparams:
                    report.err(f"Discontinued row {i}: slot {n} is empty but recommend_params_{n} is set: {rparams!r}")
                if empty_seen_at is None:
                    empty_seen_at = n
                continue
            if not (rpn and rtype):
                report.err(f"Discontinued row {i}: slot {n} is partially filled (part_number={rpn!r}, type={rtype!r})")
            if empty_seen_at is not None:
                report.err(f"Discontinued row {i}: slot {n} is filled after empty slot {empty_seen_at} (slots must fill left to right)")
            recs.append(rpn)
            rec = by_pn.get(rpn)
            if rec is None:
                report.err(f"Discontinued row {i}: recommend_part_number_{n} {rpn!r} not found in the Input CSV")
                continue
            if rpn.startswith(DDB_PREFIX):
                report.err(f"Discontinued row {i}: recommend_part_number_{n} {rpn!r} is a DDB row and must not be recommended")
            if rtype != rec.get("part_number_type", ""):
                report.err(f"Discontinued row {i}: recommend_part_number_type_{n} {rtype!r} != GDB {rec.get('part_number_type')!r} for {rpn!r}")
            if rparams != rec.get("params", ""):
                report.err(f"Discontinued row {i}: recommend_params_{n} {rparams!r} != GDB params {rec.get('params', '')!r} for {rpn!r}")
        dup_recs = [p for p, c in Counter(r for r in recs if r).items() if c > 1]
        if dup_recs:
            report.err(f"Discontinued row {i}: recommended part number(s) repeated within the row: {dup_recs}")

    for pn, c in seen.items():
        if c > 1:
            report.err(f"Discontinued file lists {pn!r} {c} times (one row per DDB part number)")
    gdb_ddb = {pn for pn in by_pn if pn and pn.startswith(DDB_PREFIX)}
    missing = gdb_ddb - set(seen)
    if missing:
        report.err(f"Discontinued file is missing row(s) for DDB part number(s) in the Input CSV: {sorted(missing)}")


def main():
    # messages contain icons (①...) / Japanese; a piped Windows stdout defaults to cp1252
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) not in (4, 5):
        print("Usage: python validate_files.py <input_csv> <spec_grouping_xlsx> <type_grouping_xlsx> [<discontinued_csv>]")
        sys.exit(2)
    input_csv, spec_xlsx, type_xlsx = sys.argv[1:4]
    report = Report()
    ctx = validate_input_csv(input_csv, report)
    validate_xlsx_pair(spec_xlsx, type_xlsx, ctx, report)
    if len(sys.argv) == 5:
        validate_discontinued(sys.argv[4], ctx["rows"], report)
    report.print_all()
    sys.exit(0 if report.ok() else 1)


if __name__ == "__main__":
    main()
